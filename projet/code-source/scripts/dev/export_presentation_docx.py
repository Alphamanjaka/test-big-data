"""Exporte le support de soutenance au format Word : une page par slide, avec l'image de la
slide, sa durée et le texte prévu à l'oral.

Usage :
    python scripts/dev/export_presentation_docx.py [--out FICHIER.docx] [--version 1]
                                                   [--slides-dir DOSSIER]

Le document est tiré du deck et du script oral, pour ne jamais en diverger :
- le deck `documents/slide_soutenance/Soutenance_M2_MBDS_RANOMENJANAHARY.pptx` donne le titre
  et les notes de présentation (texte à dire) de chaque slide, ainsi que ses images :
  PowerPoint les exporte (COM, Windows) depuis une copie, le deck lui-même n'est jamais ouvert.
  `--slides-dir` réutilise des images déjà exportées (`slide-01.png`, `slide-02.png`…) ;
- le script oral `documents/soutenance_script_oral.md` donne la durée et les repères de chaque
  slide, et le chronométrage par partie.

La slide de démonstration n'est pas reproduite : dans le deck, elle attend encore la vidéo.
Elle est remplacée par le scénario de démonstration raconté (SCENARIO), dont chaque étape
s'appuie sur une capture réelle du 30/09/2026 (`documents/captures/`, annexe H du mémoire).
"""

from __future__ import annotations

import argparse
import datetime
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation

from export_memoire_docx import add_field, set_page, shade_cell

ROOT = Path(__file__).resolve().parents[4]
DECK = ROOT / "documents" / "slide_soutenance" / "Soutenance_M2_MBDS_RANOMENJANAHARY.pptx"
SCRIPT_ORAL = ROOT / "documents" / "soutenance_script_oral.md"
CAPTURES = ROOT / "documents" / "captures"
IMAGES = ROOT / "documents" / "image"
OUT = ROOT / "documents" / "Presentation_soutenance_RANOMENJANAHARY.docx"

TITLE = "Plateforme Big Data de gestion et de gouvernance des données patients"
SUBTITLE = "Nettoyage, déduplication et contrôle d'accès basé sur le consentement"
AUTHOR = "RANOMENJANAHARY Manjaka Alpha"
SUPERVISORS = [
    ("Encadrant professionnel", "M. Harena Ny Aina RABEMANOELA (MMT)"),
    ("Encadrant pédagogique", "M. Rojo RABENANAHARY (MBDS)"),
]
INSTITUTION_LOGOS = ["logo-ituniversity.png", "logo-mbds.jpg", "logo-uca.png"]
MMT_LOGO = "mmt-logo.png"
DEMO_SLIDE = 17
MONTHS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
          "septembre", "octobre", "novembre", "décembre"]

# Couleurs et mesures du deck (build_soutenance_deck.py) ; A4 portrait aux marges du mémoire.
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x55, 0x97)
AMBER = RGBColor(0xC7, 0x84, 0x00)
MUTED = RGBColor(0x5B, 0x64, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX, TINT_HEX, AMBER_T_HEX, LINE_HEX = "1F3864", "F1F4F9", "FBF1DC", "C9D2E0"
TEXT_WIDTH_CM = 16.0
EMU_PER_INCH = 914400

# Scénario de la démonstration (slide 17). Chaque étape : ce qui est à l'écran, ce qui est dit,
# et la capture du 30/09/2026 qui en est la sortie réelle : fichier, bande verticale retenue
# (fraction de la hauteur, None = image entière) et numéro de figure dans le mémoire.
# Les chiffres cités sont ceux des captures ; la somme des durées égale celle de la slide.
SCENARIO_TITLE = "Démonstration : un patient, de la source à l'accès contrôlé"
SCENARIO_OPENING = ("0:10", (
    "Je vous propose de suivre un patient fictif, Thibaut Aubry. Comme Jean Rakoto au début, "
    "il a été enregistré par trois services : la consultation, l'imagerie et la pharmacie."))
SCENARIO = [
    {
        "titre": "Le pipeline a tourné",
        "duree": "0:30",
        "ecran": "terminal de la VM : journal du run complet, lancé avant la séance (1 min 30 s)",
        "dire": "Ses trois fiches arrivent brutes dans la zone RAW. Elles sont harmonisées dans "
                "SILVER, où Spark applique la règle d'identité stricte, puis GOLD calcule les "
                "événements et les consentements. Sur le jeu difficile : 1 057 fiches, 942 patients "
                "maîtres, 115 doublons rattachés, en une minute trente.",
        "capture": ("C12_run_pipeline.png", None, 15),
    },
    {
        "titre": "Le tableau de bord",
        "duree": "0:30",
        "ecran": "interface web, page « Tableau de bord »",
        "dire": "Le tableau de bord confirme le run : cinq étapes sur cinq, les trois zones du lac "
                "à jour. L'historique garde chaque exécution avec ses chiffres : 1 057 lignes, "
                "942 patients maîtres, 115 doublons.",
        "capture": ("C07_pipeline.png", (0.0, 0.685), 11),
    },
    {
        "titre": "Le dossier du patient",
        "duree": "0:40",
        "ecran": "interface web, fiche de Thibaut Aubry",
        "dire": "Voici notre patient : un seul identifiant pour ses trois fiches. La fiche de la "
                "pharmacie a créé le patient maître ; celles de la consultation et de l'imagerie "
                "lui sont rattachées par correspondance exacte, car le CIN, le genre, la date et la "
                "ville de naissance sont identiques. En bas, ses consentements, finalité par "
                "finalité. C'est la réponse à la première question : c'est bien lui.",
        "capture": ("C09_fiche_patient.png", None, 12),
    },
    {
        "titre": "Qui peut lire ?",
        "duree": "1:00",
        "ecran": "appels à l'API de gouvernance avec les clés de trois rôles (clés masquées)",
        "dire": "Deuxième question : qui a le droit de lire ? Sans clé, l'API répond 401. Sans "
                "finalité, ou avec une finalité inconnue comme « marketing », 422. Un lecteur qui "
                "demande le journal d'audit : 403, son rôle ne suffit pas. Un analyste qui demande "
                "les patients pour des statistiques en obtient 395 : les 547 autres n'ont pas "
                "accordé cet usage et sont filtrés. S'il vise l'un d'eux directement : 403. Enfin, "
                "pour Thibaut Aubry, qui a accepté, l'accès est accordé, avec ses trois fiches.",
        "capture": ("C14_refus_403.png", (0.0, 0.655), 16),
    },
    {
        "titre": "La trace",
        "duree": "0:20",
        "ecran": "journal d'audit des sept appels",
        "dire": "Chaque appel, accepté ou refusé, a laissé une ligne : l'utilisateur, le point "
                "d'accès, le statut, la finalité et, quand le consentement manque, le motif du refus.",
        "capture": ("C14_refus_403.png", (0.655, 1.0), 16),
    },
    {
        "titre": "La preuve",
        "duree": "0:20",
        "ecran": "évaluation du run sur la vérité terrain",
        "dire": "Dernier écran : ces 942 patients sont confrontés à la vérité terrain. Précision "
                "1,000 : aucune fusion à tort. Rappel 0,179 : le prix de la règle stricte. Et "
                "Spark donne exactement les mêmes identifiants que la référence Python.",
        "capture": ("C16_evaluation.png", None, 18),
    },
]
SCENARIO_PREPARATION = [
    ("Environnement", "VM démarrée (HDFS, Hive, Spark) ; API de gouvernance et interface web "
                      "lancées sur le poste."),
    ("Données", "jeu d'évaluation difficile, synthétique : 1 057 fiches issues de trois sources ; "
                "consentements de démonstration (2 826 avis)."),
    ("Avant la séance", "lancer le run complet (1 min 30 s) ; ouvrir les pages dans l'ordre du "
                        "scénario ; préparer les clés d'API des trois rôles, jamais affichées."),
    ("Plan de repli", "si la VM ou le poste fait défaut, dérouler le même scénario sur les "
                      "captures ci-dessus."),
    ("À dire comme tel", "données fictives ; consentements de démonstration, non recueillis "
                         "auprès de patients ; planification automatique non activée."),
]

PPT_EXPORT = r"""
$ErrorActionPreference = 'Stop'
$app = New-Object -ComObject PowerPoint.Application
try {
  $deck = $app.Presentations.Open('__SRC__', -1, 0, 0)
  foreach ($s in $deck.Slides) {
    $s.Export(('__DIR__\slide-{0:D2}.png' -f $s.SlideIndex), 'PNG', 1600, 900)
  }
  $deck.Close()
} finally {
  if ($app.Presentations.Count -eq 0) { $app.Quit() }
}
"""


# --------------------------------------------------------------------------- sources

def seconds(mark: str) -> int:
    minutes, secs = mark.split(":")
    return int(minutes) * 60 + int(secs)


def mark(total: int) -> str:
    return "{0}:{1:02d}".format(total // 60, total % 60)


def read_script_oral() -> "tuple[dict, list]":
    """Durée et repères par slide, et chronométrage par partie, lus dans le script oral."""
    text = SCRIPT_ORAL.read_text(encoding="utf-8")
    timings = {}
    for m in re.finditer(r"^## S(\d+)\. (.+?) — (\d+:\d\d)[^·\n]*· `\[(\d+:\d\d)\] → \[(\d+:\d\d)\]`",
                         text, re.M):
        timings[int(m.group(1))] = {"nom": m.group(2), "duree": m.group(3),
                                    "debut": m.group(4), "fin": m.group(5)}
    parts = [(m.group(2), m.group(3), m.group(1))
             for m in re.finditer(r"^\| `\[(\d+:\d\d)\]` \| (.+?) \| (S\d+(?: à S\d+)?) \|", text, re.M)]
    return timings, parts


def read_deck(deck: Path) -> list:
    """Surtitre, titre et notes de chaque slide.

    Le surtitre et le titre sont les zones placées par le gabarit du deck (0,86 et 1,08 pouce
    du haut) ; la page de titre et la page de remerciements n'en ont pas : leur titre est le
    texte de plus grand corps.
    """
    slides = []
    for slide in Presentation(str(deck)).slides:
        kicker = title = ""
        largest = (0, "")
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            top = shape.top / EMU_PER_INCH
            text = " ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if abs(top - 0.86) < 0.05:
                kicker = text
            elif abs(top - 1.08) < 0.05:
                title = text
            sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
            if sizes and max(sizes) > largest[0]:
                largest = (max(sizes), text)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        slides.append({"surtitre": kicker, "titre": title or largest[1], "notes": notes})
    return slides


def export_slides(deck: Path, out_dir: Path) -> None:
    """Images PNG des slides, exportées par PowerPoint depuis une copie du deck."""
    copy = out_dir / "deck.pptx"
    shutil.copyfile(str(deck), str(copy))
    script = PPT_EXPORT.replace("__SRC__", str(copy)).replace("__DIR__", str(out_dir))
    res = subprocess.run(["powershell", "-NoProfile", "-Command", script],
                         capture_output=True, text=True)
    if res.returncode != 0:
        raise SystemExit("Export des slides par PowerPoint impossible :\n{0}\nExporter les slides "
                         "en PNG (slide-01.png…) puis relancer avec --slides-dir."
                         .format(res.stderr.strip()))


# --------------------------------------------------------------------------- mise en forme

def set_fonts(doc: Document) -> None:
    """Calibri partout, comme le deck ; titres de niveau 1 en bleu marine."""
    for name in ("Normal", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        rfonts.set(qn("w:eastAsia"), "Calibri")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 17), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(6 if name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(4)


def styled(paragraph, text: str, size: float = 11, bold: bool = False, italic: bool = False,
           color: RGBColor = None):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def line(doc: Document, text: str = "", size: float = 11, bold: bool = False,
         italic: bool = False, color: RGBColor = None, align=None, after: float = 6,
         keep: bool = False):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = keep
    if text:
        styled(paragraph, text, size, bold, italic, color)
    return paragraph


# Ordre imposé par le schéma aux éléments qui suivent w:pBdr dans w:pPr : Word refuse
# d'ouvrir un fichier dont la bordure serait placée après l'alignement ou l'espacement.
PPR_AFTER_BORDER = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def frame(paragraph, color: str = LINE_HEX) -> None:
    """Filet fin autour d'un paragraphe (image de slide ou capture)."""
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:{0}".format(side))
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "1")
        edge.set(qn("w:color"), color)
        borders.append(edge)
    paragraph._p.get_or_add_pPr().insert_element_before(borders, *PPR_AFTER_BORDER)


def add_image(doc: Document, path: Path, max_width: float, max_height: float,
              band=None, framed: bool = True, keep: bool = False):
    """Image centrée tenant dans max_width × max_height (cm).

    `band` = (haut, bas) en fraction de la hauteur : seule cette bande est affichée, par un
    recadrage Word (a:srcRect) ; l'image d'origine reste entière dans le fichier.
    `keep` garde l'image sur la même page que le paragraphe suivant (sa légende).
    """
    top, bottom = band or (0.0, 1.0)
    info = DocxImage.from_file(str(path))
    width_px, height_px = info.px_width, info.px_height * (bottom - top)
    width = min(max_width, max_height * width_px / height_px)
    height = width * height_px / width_px
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = keep
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width), height=Cm(height))
    if band:
        blip = run._r.find(".//" + qn("a:blip"))
        crop = OxmlElement("a:srcRect")
        crop.set("t", str(int(round(top * 100000))))
        crop.set("b", str(int(round((1 - bottom) * 100000))))
        blip.addnext(crop)
    if framed:
        frame(paragraph)
    return paragraph


def set_table_borders(table, color: str = LINE_HEX) -> None:
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:{0}".format(side))
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        borders.append(edge)
    table._tbl.tblPr.append(borders)


def add_table(doc: Document, header: list, rows: list, widths: list, size: float = 10):
    """Tableau à en-tête bleu marine et lignes alternées, largeurs en cm."""
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, values in enumerate([header] + list(rows)):
        for j, value in enumerate(values):
            cell = table.cell(i, j)
            cell.width = Cm(widths[j])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            if i == 0:
                shade_cell(cell, NAVY_HEX)
                styled(paragraph, value, size, bold=True, color=WHITE)
            else:
                if i % 2 == 0:
                    shade_cell(cell, TINT_HEX)
                styled(paragraph, value, size)
    return table


def add_box(doc: Document, paragraphs: list, fill: str = AMBER_T_HEX, size: float = 10.5) -> None:
    """Encadré coloré d'une cellule ; chaque paragraphe est une liste de (texte, gras)."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(TEXT_WIDTH_CM)
    shade_cell(cell, fill)
    for k, parts in enumerate(paragraphs):
        paragraph = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        for text, bold in parts:
            styled(paragraph, text, size, bold=bold)
    line(doc, after=2)


# --------------------------------------------------------------------------- pages

def add_cover(doc: Document, version: str, date_text: str) -> None:
    logos = doc.add_table(rows=1, cols=2)
    logos.autofit = False
    left, right = logos.cell(0, 0), logos.cell(0, 1)
    left.width, right.width = Cm(13.5), Cm(2.5)
    paragraph = left.paragraphs[0]
    for k, name in enumerate(INSTITUTION_LOGOS):
        if k:
            paragraph.add_run("   ")
        paragraph.add_run().add_picture(str(IMAGES / name), height=Cm(0.95))
    paragraph = right.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run().add_picture(str(IMAGES / MMT_LOGO), height=Cm(1.6))

    for _ in range(4):
        line(doc)
    line(doc, "SUPPORT DE SOUTENANCE — MASTER MBDS", 12, bold=True, color=AMBER, after=10)
    line(doc, TITLE, 26, bold=True, color=NAVY, after=10)
    line(doc, SUBTITLE, 14, italic=True, color=MUTED, after=36)
    paragraph = line(doc, after=4)
    styled(paragraph, "Présenté par ", 13)
    styled(paragraph, AUTHOR, 13, bold=True)
    for label, name in SUPERVISORS:
        line(doc, "{0} : {1}".format(label, name), 11, color=MUTED, after=2)
    line(doc, "Madagascar Medical Technology — Département Recherche et Développement", 11,
         color=MUTED, after=48)
    add_box(doc, [
        [("Version {0} — document de travail du {1}".format(version, date_text), True)],
        [("Soutenance prévue en octobre 2026. Le document reprend les 20 slides du support, avec "
          "le texte prévu à l'oral ; la démonstration y est racontée sous forme de scénario.", False)],
        [("Toutes les données présentées sont synthétiques.", False)],
    ], fill=TINT_HEX, size=11)


def add_reading_notes(doc: Document, slides: list, timings: dict, parts: list) -> None:
    heading = doc.add_heading("À propos de ce document", level=1)
    heading.paragraph_format.page_break_before = True
    line(doc, "Ce document reprend, slide par slide, le support de la soutenance : {0} slides, pour "
              "un exposé qui se termine à {1}, dans les 20 minutes prévues. Chaque page montre la "
              "slide, sa durée et le texte prévu à l'oral ; la flèche (→) marque la phrase de "
              "transition vers la slide suivante.".format(len(slides), parts[-1][2] if parts else "?"))
    line(doc, "La slide {0}, consacrée à la démonstration, est rédigée sous forme de scénario : "
              "l'enchaînement des écrans prévus et ce qui sera dit pendant chacun. La vidéo n'est "
              "pas encore insérée dans le support ; chaque étape du scénario a toutefois été "
              "exécutée le 30/09/2026 sur la plateforme, et les captures qui l'illustrent en sont "
              "les sorties réelles.".format(DEMO_SLIDE))

    doc.add_heading("Déroulé", level=2)
    rows, start = [], 0
    for label, slide_range, end in parts:
        rows.append([label[0].upper() + label[1:], slide_range, mark(seconds(end) - start),
                     "{0} → {1}".format(mark(start), end)])
        start = seconds(end)
    add_table(doc, ["Partie", "Slides", "Durée", "Repères"], rows, [7.6, 2.8, 2.0, 3.6])
    line(doc, after=2)

    doc.add_heading("Liste des slides", level=2)
    rows = []
    for n, slide in enumerate(slides, 1):
        title = SCENARIO_TITLE if n == DEMO_SLIDE else slide["titre"]
        rows.append([str(n), title, timings.get(n, {}).get("duree", "")])
    add_table(doc, ["N°", "Titre", "Durée"], rows, [1.2, 12.8, 2.0], size=9.5)


def add_slide_header(doc: Document, n: int, kicker: str, title: str, timing: dict) -> None:
    paragraph = line(doc, after=2, keep=True)
    paragraph.paragraph_format.page_break_before = True
    styled(paragraph, "SLIDE {0}".format(n), 9.5, bold=True, color=AMBER)
    if kicker:
        styled(paragraph, "  ·  {0}".format(kicker.upper()), 9.5, bold=True, color=AMBER)
    heading = doc.add_heading(title, level=1)
    heading.paragraph_format.keep_with_next = True
    if timing:
        line(doc, "Durée : {0}  ·  de {1} à {2}".format(timing["duree"], timing["debut"], timing["fin"]),
             9.5, color=MUTED, after=8, keep=True)


def add_notes(doc: Document, notes: str) -> None:
    label = line(doc, "À L'ORAL", 9.5, bold=True, color=BLUE, after=3, keep=True)
    label.paragraph_format.space_before = Pt(10)
    for block in re.split(r"\n\s*\n", notes.strip()):
        text = " ".join(block.split())
        if not text:
            continue
        if text.startswith("→"):
            line(doc, text, 10.5, italic=True, color=MUTED)
        else:
            line(doc, text)


def add_slide(doc: Document, n: int, slide: dict, image: Path, timing: dict) -> None:
    add_slide_header(doc, n, slide["surtitre"], slide["titre"], timing)
    add_image(doc, image, TEXT_WIDTH_CM, 9.0)
    add_notes(doc, slide["notes"])


def add_demo(doc: Document, n: int, slide: dict, timing: dict) -> None:
    """Slide de démonstration : le scénario raconté, étape par étape, avec ses captures."""
    add_slide_header(doc, n, "Démonstration — scénario", SCENARIO_TITLE, timing)
    add_box(doc, [[("Scénario proposé. ", True),
                   ("La vidéo n'est pas encore insérée dans le support : la démonstration est "
                    "racontée ici, écran par écran. Chaque étape a été exécutée le 30/09/2026 sur "
                    "la plateforme, avec les données synthétiques du jeu d'évaluation difficile ; "
                    "les captures en sont les sorties réelles (mémoire, annexe H).", False)]])

    doc.add_heading("Sur la slide : le fil du scénario", level=2)
    rows = [["", "Ouverture : le patient suivi", "la slide", SCENARIO_OPENING[0]]]
    rows += [[str(k), step["titre"], step["ecran"], step["duree"]] for k, step in enumerate(SCENARIO, 1)]
    rows.append(["", "Total", "", timing["duree"] if timing else ""])
    add_table(doc, ["", "Étape", "À l'écran", "Durée"], rows, [0.8, 4.4, 9.2, 1.6])

    doc.add_heading("Déroulé détaillé", level=2)
    paragraph = line(doc, after=4)
    styled(paragraph, "Ouverture ({0}). ".format(SCENARIO_OPENING[0]), bold=True, color=NAVY)
    styled(paragraph, "« {0} »".format(SCENARIO_OPENING[1]))
    for k, step in enumerate(SCENARIO, 1):
        paragraph = line(doc, after=2, keep=True)
        paragraph.paragraph_format.space_before = Pt(8)
        styled(paragraph, "{0}. {1} ({2})".format(k, step["titre"], step["duree"]), 11.5,
               bold=True, color=NAVY)
        paragraph = line(doc, after=3, keep=True)
        styled(paragraph, "À l'écran : ", 10.5, bold=True, color=MUTED)
        styled(paragraph, step["ecran"] + ".", 10.5, color=MUTED)
        paragraph = line(doc, after=4, keep=True)
        styled(paragraph, "« {0} »".format(step["dire"]))
        name, band, figure = step["capture"]
        path = CAPTURES / name
        if path.exists():
            add_image(doc, path, TEXT_WIDTH_CM, 11.5, band, keep=True)
            line(doc, "Capture du 30/09/2026, sortie réelle, clés masquées — mémoire, annexe H, "
                      "figure {0}.".format(figure), 9, italic=True, color=MUTED,
                 align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
        else:
            print("  [avertissement] capture introuvable : {0}".format(path))

    doc.add_heading("Préparation et plan de repli", level=2)
    add_table(doc, ["", "Détail"], [[a, b] for a, b in SCENARIO_PREPARATION], [3.6, 12.4])
    transitions = [b for b in re.split(r"\n\s*\n", slide["notes"].strip()) if b.strip().startswith("→")]
    if transitions:
        line(doc, after=2)
        line(doc, " ".join(transitions[-1].split()), 10.5, italic=True, color=MUTED)


def add_footer(doc: Document, version: str, date_text: str) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styled(paragraph, "Support de soutenance — version {0} du {1} — page ".format(version, date_text),
           9, color=MUTED)
    add_field(paragraph, "PAGE")
    for run in paragraph.runs:
        run.font.size = Pt(9)


# --------------------------------------------------------------------------- assemblage

def build(out: Path, version: str, slides_dir: Path = None) -> None:
    timings, parts = read_script_oral()
    slides = read_deck(DECK)
    demo = sum(seconds(step["duree"]) for step in SCENARIO) + seconds(SCENARIO_OPENING[0])
    if DEMO_SLIDE in timings and demo != seconds(timings[DEMO_SLIDE]["duree"]):
        raise SystemExit("Scénario de {0} pour une slide de {1} : ajuster les durées de SCENARIO."
                         .format(mark(demo), timings[DEMO_SLIDE]["duree"]))
    today = datetime.date.today()
    date_text = "{0} {1} {2}".format(today.day, MONTHS[today.month - 1], today.year)

    with tempfile.TemporaryDirectory() as tmp:
        if slides_dir is None:
            slides_dir = Path(tmp)
            export_slides(DECK, slides_dir)
        images = sorted(slides_dir.glob("slide-*.png"))
        if len(images) != len(slides):
            raise SystemExit("{0} images pour {1} slides dans {2}.".format(len(images), len(slides), slides_dir))

        doc = Document()
        set_page(doc.sections[0], landscape=False)
        set_fonts(doc)
        doc.core_properties.title = "Support de soutenance — version {0}".format(version)
        doc.core_properties.author = AUTHOR
        add_cover(doc, version, date_text)
        add_reading_notes(doc, slides, timings, parts)
        for n, (slide, image) in enumerate(zip(slides, images), 1):
            if n == DEMO_SLIDE:
                add_demo(doc, n, slide, timings.get(n))
            else:
                add_slide(doc, n, slide, image, timings.get(n))
        add_footer(doc, version, date_text)
        out.parent.mkdir(parents=True, exist_ok=True)
        try:
            doc.save(str(out))
        except PermissionError:
            raise SystemExit("Impossible d'écrire {0} : le fichier est probablement ouvert dans Word."
                             .format(out.name))
    print("DOCX écrit : {0}".format(out.resolve()))
    print("Slides : {0} (dont la démonstration en scénario : {1} étapes, {2}) · version {3} du {4}"
          .format(len(slides), len(SCENARIO), mark(demo), version, date_text))


def main() -> None:
    parser = argparse.ArgumentParser(description="Exporte le support de soutenance en .docx.")
    parser.add_argument("--out", type=Path, default=OUT, help="chemin du fichier .docx")
    parser.add_argument("--version", default="1", help="numéro de version affiché")
    parser.add_argument("--slides-dir", type=Path, default=None,
                        help="dossier d'images déjà exportées (slide-01.png…)")
    args = parser.parse_args()
    build(args.out, args.version, args.slides_dir)


if __name__ == "__main__":
    main()
