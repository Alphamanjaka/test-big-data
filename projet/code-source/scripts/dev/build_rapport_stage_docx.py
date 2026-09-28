"""Construit le rapport de stage (.docx) sur le gabarit Word des rapports de référence MBDS.

Usage :
    python scripts/dev/build_rapport_stage_docx.py [--source ...] [--out ...] [--no-word]

Le document part du rapport de référence `documents/Rapport de stage ETU 1156 ... .docx` utilisé
comme **gabarit** : ses styles (Times New Roman, titres numérotés 1 / 1.1 / 1.1.1, légendes,
tables des matières), sa page de garde à quatre logos (UCA, MBDS, MMT, IT University), ses
en-têtes et pieds de page et son découpage en sections (garde, liminaires en chiffres romains,
corps en chiffres arabes) sont conservés. Tout le contenu du gabarit est retiré, y compris les
images qu'il embarquait : seul le texte de `documents/rapport_stage_source.md` est injecté.

Syntaxe reconnue dans la source (sous-ensemble de Markdown) :
    # Titre            titre de chapitre numéroté (Titre 1)
    #! Titre           titre de niveau 1 non numéroté (introduction, conclusion, listes...)
    #= Titre           titre liminaire hors sommaire (remerciements, résumé, abstract)
    ## / ###           titres numérotés de niveau 2 / 3 ; `##! Titre` non numéroté
    Tableau: légende   légende du tableau Markdown qui suit (numérotée par un champ SEQ)
    Figure: légende | chemin | largeur_cm [| paysage]
    Capture: C01 | légende | fichier.png | consigne
                       capture fournie par l'auteur dans documents/captures/ ; tant que le
                       fichier manque, un cadre « EMPLACEMENT RÉSERVÉ » affiche la consigne
    Code: X01 | légende | chemin::fonction[,autre] (ou chemin:début-fin) | [image.png]
                       extrait du code réel du dépôt (projet/code-source), numéros de ligne
                       compris ; remplacé par l'image si elle existe dans documents/captures/
    > texte            encadré de note
    - / 1.             listes à puces / numérotées
    [[TOC]] [[LOT]] [[LOF]] [[LOC]]  sommaire, listes des tableaux, figures, extraits de code
    [[FIN_LIMINAIRES]] fin de la section liminaire (la numérotation arabe repart à 1)
    [[SAUT]]           saut de page
Dans un tableau dont la légende contient `{gantt}`, les cellules ■ / □ / ○ deviennent des
cases colorées (réalisé daté / déclaré / prévu).

Word (COM) met ensuite à jour les champs (sommaire, listes, numéros) : sans Word, passer
`--no-word` et appuyer sur F9 à l'ouverture du document.
"""

from __future__ import annotations

import argparse
import copy
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.image.image import Image as DocxImage
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from lxml.etree import tostring as _tostring


def etree_tostring(el):
    return _tostring(el, encoding="unicode")

ROOT = Path(__file__).resolve().parents[4]
TEMPLATE = ROOT / "documents" / "Rapport de stage ETU 1156 RAMANANTSAFIDY Jonah Fitia.docx"
SOURCE = ROOT / "documents" / "rapport_stage_source.md"
OUT = ROOT / "documents" / "Rapport_de_stage_RANOMENJANAHARY_Manjaka_Alpha.docx"

META = {
    "titre": ["Plateforme Big Data de gestion et de", "gouvernance des données patients"],
    "sous_titre": "Nettoyage, déduplication et contrôle d'accès basé sur le consentement",
    "auteur": "RANOMENJANAHARY Manjaka Alpha",
    "date": "Octobre, 2026",
    "encadrants": [
        "M. Harena Ny Aina RABEMANOELA, encadrant professionnel (MMT)",
        "M. Rojo RABENANAHARY, encadrant pédagogique (MBDS)",
    ],
    "copyright": "© RANOMENJANAHARY Manjaka Alpha, 2026",
}

TEXT_WIDTH_CM = 16.0          # A4 moins deux marges de 2,5 cm
LANDSCAPE_WIDTH_CM = 24.7     # 29,7 moins deux marges de 2,5 cm
LANDSCAPE_HEIGHT_CM = 13.5    # 21 moins les marges, la légende et le pied de page
PORTRAIT_FIG_HEIGHT_CM = 19.0
KEEP_TOGETHER_ROWS = 14       # au-delà, un tableau peut se poursuivre page suivante
CAPTURES_ROOT = ROOT / "documents" / "captures"
CODE_ROOT = ROOT / "projet" / "code-source"
PLACEHOLDER_FILL = "F4F6FA"
PLACEHOLDER_HEIGHT_CM = 5.5
CODE_FILL = "F6F8FA"
CODE_PT = 8.5
LOGO_DIR = ROOT / "documents" / "figures" / "logos"
LOGO_WIDTH_CM = 2.3           # logos normalisés 400 × 140 : ≈ 0,8 cm de haut
LOGO_TABLE_WIDTHS = [3.3, 2.7, 1.9]   # Usage, Outil, Version (cm) ; Description = le reste
ACCENT = RGBColor(0x1F, 0x38, 0x64)
HEADER_FILL = "DCE3EE"
NOTE_FILL = "F2F5F9"
GANTT_FILLS = {"■": "1F3864", "□": "8EAADB", "○": "D9D9D9"}
BULLET_NUM_ID = 90
ORDERED_ABSTRACT_ID = 91


# --------------------------------------------------------------------------- XML utilitaires

def _set(el, tag, **attrs):
    """Retourne l'enfant `tag` de `el` (créé au besoin) avec les attributs donnés."""
    child = el.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        el.append(child)
    for key, value in attrs.items():
        child.set(qn(key), str(value))
    return child


def shade(cell_or_ppr, fill):
    tc_pr = cell_or_ppr
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_field(paragraph, instr, placeholder=""):
    """Champ complexe (begin / instr / separate / résultat / end), mis à jour par Word."""
    def run_with(child):
        r = OxmlElement("w:r")
        r.append(child)
        paragraph._p.append(r)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    t = OxmlElement("w:t")
    t.text = placeholder
    run_with(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def no_numbering(paragraph):
    num_pr = _set(paragraph._p.get_or_add_pPr(), "w:numPr")
    _set(num_pr, "w:ilvl", **{"w:val": 0})
    _set(num_pr, "w:numId", **{"w:val": 0})


# --------------------------------------------------------------------------- texte en ligne

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(paragraph, text, size=None, bold=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 12) - 2)
        elif part.startswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if bold:
            run.bold = True
        if size and not part.startswith("`"):
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


# --------------------------------------------------------------------------- constructeur

class Builder:
    def __init__(self, template: Path):
        self.doc = Document(str(template))
        self.body = self.doc.element.body
        self.counters = {"Tableau": 0, "Figure": 0, "Extrait": 0}
        self.slots = []               # (type, identifiant, fichier attendu, présent ?)
        self.pending_caption = None
        self.body_started = False     # la première section du corps repart à 1
        self._ordered_lists = 0
        self._prepare_template()

    # ---- gabarit -------------------------------------------------------------------------
    def _prepare_template(self):
        body = self.body
        sect_paras = [p for p in body.iterchildren(qn("w:p"))
                      if p.find(qn("w:pPr")) is not None
                      and p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None]
        cover_end, front_end = sect_paras[0], sect_paras[1]
        self.front_sect = copy.deepcopy(front_end.find(qn("w:pPr")).find(qn("w:sectPr")))
        final = body.find(qn("w:sectPr"))
        self.body_sect = copy.deepcopy(final)

        # Tout retirer après la page de garde, sauf la dernière section.
        remove = False
        for child in list(body):
            if remove and child is not final:
                body.remove(child)
            if child is cover_end:
                remove = True

        # Section liminaire : marges symétriques, numérotation romaine à partir de i.
        mar = self.front_sect.find(qn("w:pgMar"))
        mar.set(qn("w:right"), "1418")
        # Section du corps : chaque page numérotée (pas de première page sans numéro).
        for el in list(self.body_sect):
            if el.tag == qn("w:titlePg") or (
                    el.tag == qn("w:footerReference") and el.get(qn("w:type")) == "first"):
                self.body_sect.remove(el)
        for el in list(final):
            final.remove(el)
        for el in self.body_sect:
            final.append(copy.deepcopy(el))
        self.final_sect = final

        self._fill_cover()
        self._tune_styles()
        self._add_list_definitions()

    def _cover_paragraph(self, needle):
        for p in self.doc.paragraphs:
            # Les zones à remplir du gabarit sont des champs MACROBUTTON : leur texte est
            # dans le code du champ, pas dans le texte affiché du paragraphe.
            if needle in p.text or needle in p._p.xml:
                return p
        raise SystemExit(f"Page de garde du gabarit : paragraphe « {needle} » introuvable.")

    @staticmethod
    def _replace_text(paragraph, lines, bold=None, size=None):
        p = paragraph._p
        for child in list(p):
            if child.tag != qn("w:pPr"):
                p.remove(child)
        for i, line in enumerate(lines):
            run = paragraph.add_run(line)
            if bold is not None:
                run.bold = bold
            if size:
                run.font.size = Pt(size)
            if i < len(lines) - 1:
                run.add_break()

    def _fill_cover(self):
        title1 = self._cover_paragraph("Data Lake interopérable")
        title2 = self._cover_paragraph("Une application web de visualisation")
        subtitle = self._cover_paragraph("Tapez ici le sous-titre")
        self._replace_text(title1, [META["titre"][0]], bold=True, size=16)
        self._replace_text(title2, [META["titre"][1]], bold=True, size=16)
        self._replace_text(subtitle, [META["sous_titre"]], bold=False, size=13)
        for r in subtitle.runs:
            r.italic = True
        self._replace_text(self._cover_paragraph("RAMANANTSAFIDY Mariella"), [META["auteur"]], bold=True)
        self._replace_text(self._cover_paragraph("Octobre, 2025"), [META["date"]])
        jury = self._cover_paragraph("Jury")
        self._replace_text(jury, ["Encadrants"] + META["encadrants"])
        jury.alignment = WD_ALIGN_PARAGRAPH.CENTER
        jury.paragraph_format.first_line_indent = Cm(0)
        jury.runs[0].bold = True
        self._replace_text(self._cover_paragraph("©"), [META["copyright"]])
        props = self.doc.core_properties
        props.title = " ".join(META["titre"])
        props.subject = META["sous_titre"]
        props.author = META["auteur"]
        props.last_modified_by = META["auteur"]
        props.keywords = "Big Data, Medallion, Spark, Hive, FHIR, déduplication, MPI, consentement"
        props.comments = ""
        props.category = "Rapport de stage M2 MBDS"

    def _tune_styles(self):
        st = self.doc.styles
        para = st["Para"]
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13)):
            s = st[name]
            s.font.color.rgb = ACCENT
            s.font.size = Pt(size)
            s.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 12)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True
        st["Heading 1"].paragraph_format.page_break_before = True
        for name in ("1|Resume", "1|TableListe", "1|TitreFront"):
            st[name].font.color.rgb = ACCENT
            st[name].paragraph_format.space_after = Pt(12)
        cap = st["Caption"]
        cap.font.size = Pt(10)
        cap.font.italic = False
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.line_spacing = 1.0

    def _add_list_definitions(self):
        numbering = self.doc.part.numbering_part.element
        bullet = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{BULLET_NUM_ID}">'
            '<w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            '<w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="567" w:hanging="283"/></w:pPr>'
            '<w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/></w:rPr></w:lvl>'
            '</w:abstractNum>')
        ordered = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{ORDERED_ABSTRACT_ID}">'
            '<w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="567" w:hanging="340"/></w:pPr></w:lvl>'
            '</w:abstractNum>')
        # Titres de niveau 3 : le gabarit les décale de 3,5 cm ; on les aligne sur les
        # niveaux 1 et 2 (numéro au bord de la marge).
        heading_abstract = None
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == "1":
                heading_abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        for absn in numbering.findall(qn("w:abstractNum")):
            if absn.get(qn("w:abstractNumId")) != heading_abstract:
                continue
            for lvl in absn.findall(qn("w:lvl")):
                if lvl.get(qn("w:ilvl")) == "2":
                    ppr = lvl.find(qn("w:pPr"))
                    for tab in ppr.iter(qn("w:tab")):
                        tab.set(qn("w:pos"), "0")
                    ind = ppr.find(qn("w:ind"))
                    ind.set(qn("w:left"), "720")
                    ind.set(qn("w:hanging"), "720")
        first_num = numbering.find(qn("w:num"))
        first_num.addprevious(bullet)
        first_num.addprevious(ordered)
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{BULLET_NUM_ID}">'
            f'<w:abstractNumId w:val="{BULLET_NUM_ID}"/></w:num>'))
        self._numbering = numbering

    def _new_ordered_num(self):
        self._ordered_lists += 1
        num_id = 100 + self._ordered_lists
        self._numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}">'
            f'<w:abstractNumId w:val="{ORDERED_ABSTRACT_ID}"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        return num_id

    def _clean_headers_footers(self):
        """Retire le marquage « Confidential » que le gabarit hérite d'une étiquette de
        sensibilité d'une autre organisation : zones de texte des pieds de page.

        On parcourt les parties du paquet et non `section.footer` : cette API crée une
        définition vide pour toute section qui hérite de la précédente, ce qui effacerait
        la numérotation des pages."""
        for part in self.doc.part.package.parts:
            if not re.match(r"^/word/(header|footer)\d+\.xml$", str(part.partname)):
                continue
            for run in list(part._element.iter(qn("w:r"))):
                if run.getparent() is not None and "Confidential" in etree_tostring(run):
                    run.getparent().remove(run)

    # ---- insertion -----------------------------------------------------------------------
    def _p(self, style="Para"):
        p = self.doc.add_paragraph(style=style)
        # add_paragraph ajoute à la fin du corps, avant la sectPr finale : c'est voulu.
        return p

    def heading(self, text, level, numbered=True):
        p = self._p(f"Heading {level}")
        add_inline(p, text)
        if not numbered:
            no_numbering(p)
        return p

    def front_title(self, text):
        p = self._p("1|Resume")
        p.paragraph_format.page_break_before = True
        p.add_run(text)

    def paragraph(self, text):
        p = self._p("Para")
        if re.match(r"^\[\d+\]\s", text):
            # Référence bibliographique : alignée à gauche (les URL ne se justifient pas),
            # retrait suspendu sous le numéro.
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
        add_inline(p, text)
        return p

    def list_item(self, text, ordered_num_id=None):
        p = self._p("Para")
        p.paragraph_format.space_after = Pt(3)
        num_pr = _set(p._p.get_or_add_pPr(), "w:numPr")
        _set(num_pr, "w:ilvl", **{"w:val": 0})
        _set(num_pr, "w:numId", **{"w:val": ordered_num_id or BULLET_NUM_ID})
        add_inline(p, text)

    def note(self, text):
        p = self._p("Para")
        ppr = p._p.get_or_add_pPr()
        bdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" '
            'w:color="1F3864"/></w:pBdr>')
        ppr.append(bdr)
        shade(ppr, NOTE_FILL)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.right_indent = Cm(0.2)
        p.paragraph_format.line_spacing = 1.15
        add_inline(p, text, size=11)

    def caption(self, kind, text, keep_next=False, label=None):
        self.counters[kind] += 1
        p = self._p("Caption")
        r = p.add_run(f"{label or kind} ")
        r.bold = True
        add_field(p, f"SEQ {kind} \\* ARABIC", str(self.counters[kind]))
        for r in p.runs:
            r.bold = True
        r = p.add_run(" : ")
        r.bold = True
        add_inline(p, text, size=10)
        p.paragraph_format.keep_with_next = keep_next
        return p

    def field_block(self, instr, placeholder):
        p = self._p("Normal")
        add_field(p, instr, placeholder)

    def page_break(self):
        self._p("Normal").add_run().add_break(WD_BREAK.PAGE)

    def end_section(self, sect_template, *, landscape=False, restart=False):
        """Termine la section courante avec une copie de `sect_template`."""
        sect = copy.deepcopy(sect_template)
        for el in list(sect):
            if el.tag == qn("w:pgNumType"):
                sect.remove(el)
        pg = OxmlElement("w:pgNumType")
        if sect_template is self.front_sect:
            pg.set(qn("w:fmt"), "lowerRoman")
            pg.set(qn("w:start"), "1")
        else:
            pg.set(qn("w:fmt"), "decimal")
            if restart:
                pg.set(qn("w:start"), "1")
        sz = sect.find(qn("w:pgSz"))
        sz.addnext(pg) if sz is not None else sect.append(pg)
        if landscape:
            w, h = sz.get(qn("w:w")), sz.get(qn("w:h"))
            sz.set(qn("w:w"), h)
            sz.set(qn("w:h"), w)
            sz.set(qn("w:orient"), "landscape")
        # Paragraphe porteur du saut de section : minuscule, pour ne jamais créer de page vide.
        p = self._p("Normal")
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(1)
        _set(_set(p._p.get_or_add_pPr(), "w:rPr"), "w:sz", **{"w:val": 2})
        p._p.get_or_add_pPr().append(sect)

    def figure(self, caption, path, width_cm, landscape=False):
        if landscape:
            self.end_section(self.body_sect, restart=not self.body_started)
            self.body_started = True
        p = self._p("Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        # Largeur demandée, bornée par la largeur de texte et par la hauteur disponible
        # (la légende doit tenir sur la même page que l'image).
        img = DocxImage.from_file(str(path))
        px_w, px_h = img.px_width, img.px_height
        max_w = LANDSCAPE_WIDTH_CM if landscape else TEXT_WIDTH_CM
        max_h = LANDSCAPE_HEIGHT_CM if landscape else PORTRAIT_FIG_HEIGHT_CM
        width = min(width_cm, max_w, max_h * px_w / px_h)
        p.add_run().add_picture(str(path), width=Cm(width))
        self.caption("Figure", caption)
        if landscape:
            self.end_section(self.body_sect, landscape=True)

    # ---- emplacements : captures et extraits de code ---------------------------------
    def _boxed_cell(self, fill, border_val="single", color="8C9BB0", accent_left=False):
        """Tableau d'une cellule pleine largeur, servant de cadre."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = table._tbl.tblPr
        _set(tbl_pr, "w:tblLayout", **{"w:type": "fixed"})
        sides = "".join(
            f'<w:{b} w:val="{border_val}" w:sz="{18 if accent_left and b == "left" else 6}" '
            f'w:space="0" w:color="{"1F3864" if accent_left and b == "left" else color}"/>'
            for b in ("top", "left", "bottom", "right"))
        tbl_pr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}>{sides}</w:tblBorders>'))
        cell = table.rows[0].cells[0]
        cell.width = Cm(TEXT_WIDTH_CM)
        shade(cell._tc.get_or_add_tcPr(), fill)
        return table, cell

    def _spacer(self):
        self._p("Normal").paragraph_format.space_before = Pt(0)

    def capture(self, caption, rel_path, instruction, slot):
        """Capture fournie par l'auteur : l'image si elle existe, sinon un cadre réservé."""
        path = CAPTURES_ROOT / rel_path
        found = path.exists()
        self.slots.append(("capture", slot, rel_path, found))
        if found:
            self.figure(caption, path, TEXT_WIDTH_CM)
            return
        table, cell = self._boxed_cell(PLACEHOLDER_FILL, border_val="dashed")
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tr_pr.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{int(PLACEHOLDER_HEIGHT_CM * 567)}" '
                               'w:hRule="atLeast"/>'))
        _set(cell._tc.get_or_add_tcPr(), "w:vAlign", **{"w:val": "center"})
        lines = [(f"EMPLACEMENT RÉSERVÉ — {slot}", 12, True, False),
                 (f"Fichier attendu : documents/captures/{rel_path}", 10, False, False),
                 (instruction, 10, False, True)]
        first = True
        for text, size, bold, italic in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.size = Pt(size)
            run.bold = bold
            run.italic = italic
            run.font.color.rgb = ACCENT if bold else RGBColor(0x59, 0x59, 0x59)
        self.caption("Figure", caption)

    def code(self, caption, spec, image_rel, slot):
        """Extrait de code : capture de l'auteur si fournie, sinon le code réel du dépôt."""
        image = CAPTURES_ROOT / image_rel if image_rel else None
        rel_file, numbered = extract_code(spec)
        where = f"{rel_file}, l. {numbered[0][0]}–{numbered[-1][0]}" if numbered else rel_file
        self.slots.append(("code", slot, image_rel, bool(image and image.exists())))
        if image and image.exists():
            self.figure(f"{caption} ({where})", image, TEXT_WIDTH_CM)
            return
        table, cell = self._boxed_cell(CODE_FILL, accent_left=True)
        first = True
        for lineno, text in numbered:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.keep_with_next = True
            # Une ligne trop longue revient à la ligne sous le code, pas sous son numéro.
            pf.left_indent = Cm(1.05)
            pf.first_line_indent = Cm(-1.05)
            num = p.add_run(f"{lineno:>4}  " if lineno else "      ")
            num.font.name = "Consolas"
            num.font.size = Pt(CODE_PT)
            num.font.color.rgb = RGBColor(0x9A, 0xA3, 0xAE)
            run = p.add_run(text.replace("\t", "    "))
            run.font.name = "Consolas"
            run.font.size = Pt(CODE_PT)
            if lineno is None:
                run.font.color.rgb = RGBColor(0x9A, 0xA3, 0xAE)
                run.italic = True
        self.caption("Extrait", f"{caption} ({where})", label="Extrait de code")

    def _logo_cell(self, cell, text):
        """Cellule `logo:<id> Nom` : le logo normalisé, puis le nom de l'outil dessous."""
        key, _, name = text[len("logo:"):].partition(" ")
        path = LOGO_DIR / f"{key}.png"
        if not path.exists():
            raise SystemExit(f"Logo absent : {path} (lancer scripts/dev/build_logos.py)")
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.add_run().add_picture(str(path), width=Cm(LOGO_WIDTH_CM))
        if name:
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.0
            run = p.add_run(name.strip())
            run.font.size = Pt(8.5)
            run.bold = True
            run.font.color.rgb = ACCENT

    def table(self, rows, aligns, caption):
        gantt = "{gantt}" in caption
        logos = "{logos}" in caption
        caption = caption.replace("{gantt}", "").replace("{logos}", "").strip()
        self.caption("Tableau", caption, keep_next=True)
        ncols = len(rows[0])
        widths = self._column_widths(rows, gantt)
        if logos:
            # Usage | Outil (logo) | Version | Description : la description prend le reste.
            widths = LOGO_TABLE_WIDTHS + [TEXT_WIDTH_CM - sum(LOGO_TABLE_WIDTHS)]
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = table._tbl.tblPr
        _set(tbl_pr, "w:tblLayout", **{"w:type": "fixed"})
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            + "".join(f'<w:{b} w:val="single" w:sz="4" w:space="0" w:color="8C9BB0"/>'
                      for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
            + "</w:tblBorders>")
        old = tbl_pr.find(qn("w:tblBorders"))
        if old is not None:
            tbl_pr.remove(old)
        tbl_pr.append(borders)
        for i, row in enumerate(rows):
            tr = table.rows[i]
            tr_pr = tr._tr.get_or_add_trPr()
            _set(tr_pr, "w:cantSplit")
            if i == 0:
                _set(tr_pr, "w:tblHeader")
            for j, text in enumerate(row):
                cell = tr.cells[j]
                cell.width = Cm(widths[j])
                tc_pr = cell._tc.get_or_add_tcPr()
                cp = cell.paragraphs[0]
                cp.style = self.doc.styles["Normal"]
                pf = cp.paragraph_format
                pf.space_before = Pt(2)
                pf.space_after = Pt(2)
                pf.line_spacing = 1.0
                # Un tableau court reste sur une seule page ; un long garde au moins son
                # en-tête avec sa première ligne.
                if (len(rows) <= KEEP_TOGETHER_ROWS and i < len(rows) - 1) or i == 0:
                    pf.keep_with_next = True
                align = aligns[j] if j < len(aligns) else "left"
                if gantt and j > 0:
                    align = "center"
                cp.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                                "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(align, WD_ALIGN_PARAGRAPH.LEFT)
                if i == 0:
                    shade(tc_pr, HEADER_FILL)
                    add_inline(cp, text, size=10, bold=True, color=ACCENT)
                elif gantt and text.strip() in GANTT_FILLS:
                    shade(tc_pr, GANTT_FILLS[text.strip()])
                elif logos:
                    _set(tc_pr, "w:vAlign", **{"w:val": "center"})
                    if text.startswith("logo:"):
                        self._logo_cell(cell, text)
                    elif text != "^":
                        add_inline(cp, text, size=10, bold=(j == 0))
                else:
                    add_inline(cp, text, size=10)
        # Un groupe (lignes suivies d'un `^`) reste sur une même page.
        if logos:
            for i in range(1, len(rows) - 1):
                if rows[i + 1][0] == "^":
                    for c in table.rows[i].cells:
                        for p in c.paragraphs:
                            p.paragraph_format.keep_with_next = True
        # `^` : la cellule prolonge celle du dessus (regroupement par usage).
        for j in range(ncols):
            top = None
            for i in range(1, len(rows)):
                if j < len(rows[i]) and rows[i][j] == "^" and top is not None:
                    top.merge(table.rows[i].cells[j])
                else:
                    top = table.rows[i].cells[j]
        # Espace après le tableau.
        self._p("Normal").paragraph_format.space_before = Pt(0)

    @staticmethod
    def _column_widths(rows, gantt):
        ncols = len(rows[0])
        if gantt:
            first = 5.6
            rest = (TEXT_WIDTH_CM - first) / (ncols - 1)
            return [first] + [rest] * (ncols - 1)
        cells = [[re.sub(r"[*`]", "", r[j]) if j < len(r) else "" for j in range(ncols)]
                 for r in rows]
        # Largeur minimale : le mot le plus long de la colonne ne doit jamais être coupé
        # (≈ 0,2 cm par caractère en 10 pt, plus les marges de cellule).
        minimum = [max(len(w) for r in cells for w in (r[j].split() or [""])) * 0.2 + 0.4
                   for j in range(ncols)]
        weights = [min(max(max(len(r[j]) for r in cells), 6), 60) ** 0.8 for j in range(ncols)]
        widths = [max(TEXT_WIDTH_CM * w / sum(weights), m) for w, m in zip(weights, minimum)]
        # Redistribuer l'excédent sur les colonnes qui dépassent leur minimum.
        excess = sum(widths) - TEXT_WIDTH_CM
        if excess > 0:
            slack = [w - m for w, m in zip(widths, minimum)]
            total_slack = sum(s for s in slack if s > 0) or 1
            widths = [w - excess * max(s, 0) / total_slack for w, s in zip(widths, slack)]
        return widths

    # ---- fin -----------------------------------------------------------------------------
    def finish(self, out: Path):
        if not self.body_started:
            pg = self.final_sect.find(qn("w:pgNumType"))
            if pg is None:
                pg = OxmlElement("w:pgNumType")
                self.final_sect.find(qn("w:pgSz")).addnext(pg)
            pg.set(qn("w:fmt"), "decimal")
            pg.set(qn("w:start"), "1")
        else:
            pg = self.final_sect.find(qn("w:pgNumType"))
            if pg is None:
                pg = OxmlElement("w:pgNumType")
                self.final_sect.find(qn("w:pgSz")).addnext(pg)
            pg.set(qn("w:fmt"), "decimal")
        self._drop_unused_parts()
        self._clean_headers_footers()
        self._clean_doc_props()
        try:
            self.doc.save(str(out))
        except PermissionError:
            raise SystemExit(f"Impossible d'écrire {out.name} : le document est probablement "
                             "ouvert dans Word. Fermez-le puis relancez la génération.")

    def _clean_doc_props(self):
        """Les propriétés du gabarit décrivent un autre document et une autre organisation
        (société, gestionnaire, étiquette de sensibilité) : on ne garde que le neutre."""
        for part in self.doc.part.package.parts:
            name = str(part.partname)
            if name == "/docProps/custom.xml":
                part._blob = (
                    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
                    b'custom-properties" xmlns:vt="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/docPropsVTypes"/>')
            elif name == "/docProps/app.xml":
                xml = part.blob.decode("utf-8")
                for tag in ("Company", "Manager", "HeadingPairs", "TitlesOfParts", "Template"):
                    xml = re.sub(rf"<{tag}>.*?</{tag}>|<{tag}/>", "", xml, flags=re.S)
                part._blob = xml.encode("utf-8")

    def _drop_unused_parts(self):
        """Retire du paquet les images et liens du gabarit que plus rien ne référence."""
        part = self.doc.part
        xml = self.doc.element.xml
        keep_types = ("styles", "numbering", "settings", "fontTable", "webSettings", "theme",
                      "footnotes", "endnotes", "header", "footer", "customXml", "comments")
        for rid, rel in list(part.rels.items()):
            if any(k in rel.reltype for k in keep_types):
                continue
            if f'"{rid}"' not in xml:
                del part.rels[rid]


# --------------------------------------------------------------------------- extraits de code

def extract_code(spec):
    """Lit un extrait dans le code du dépôt, avec ses vrais numéros de ligne.

    `chemin::fonction[,autre]` : les définitions Python nommées (docstring remplacée par
    une ligne « … ») — l'extrait suit le code même si ses lignes bougent ;
    `chemin:début-fin` : une plage de lignes ; `chemin` : le fichier entier.
    Retourne (chemin relatif, [(numéro ou None, texte)]).
    """
    import ast

    symbols, span = [], None
    if "::" in spec:
        rel, names = spec.split("::", 1)
        symbols = [n.strip() for n in names.split(",")]
    elif re.search(r":\d+-\d+$", spec):
        rel, rng = spec.rsplit(":", 1)
        span = tuple(int(x) for x in rng.split("-"))
    else:
        rel = spec
    path = CODE_ROOT / rel
    if not path.exists():
        raise SystemExit(f"Extrait de code : fichier introuvable {path}")
    lines = path.read_text(encoding="utf-8").splitlines()
    if span:
        return rel, [(n, lines[n - 1]) for n in range(span[0], span[1] + 1)]
    if not symbols:
        return rel, [(n + 1, t) for n, t in enumerate(lines)]

    tree = ast.parse("\n".join(lines))
    nodes = {n.name: n for n in ast.walk(tree)
             if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))}
    out = []
    for name in symbols:
        if name not in nodes:
            raise SystemExit(f"Extrait de code : « {name} » absent de {rel}")
        node = nodes[name]
        start = min([node.lineno] + [d.lineno for d in node.decorator_list])
        skip = set()
        first = node.body[0] if node.body else None
        if (isinstance(first, ast.Expr) and isinstance(getattr(first, "value", None), ast.Constant)
                and isinstance(first.value.value, str)):
            skip = set(range(first.lineno, first.end_lineno + 1))
        if out:
            out.append((None, ""))
        for n in range(start, node.end_lineno + 1):
            if n in skip:
                if n == first.lineno:
                    indent = re.match(r"\s*", lines[n - 1]).group(0)
                    out.append((None, f'{indent}"""…"""  (docstring omise)'))
                continue
            out.append((n, lines[n - 1]))
    return rel, out


# --------------------------------------------------------------------------- lecture source

TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def parse_aligns(sep):
    out = []
    for c in split_row(sep):
        if c.startswith(":") and c.endswith(":"):
            out.append("center")
        elif c.endswith(":"):
            out.append("right")
        else:
            out.append("left")
    return out


def build(source: Path, out: Path):
    b = Builder(TEMPLATE)
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    para_buf: list[str] = []
    ordered_id = None

    def flush():
        if para_buf:
            b.paragraph(" ".join(s.strip() for s in para_buf))
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not re.match(r"^\d+\.\s", s):
            ordered_id = None    # toute autre ligne clôt la liste numérotée en cours
        if s.startswith("<!--"):
            flush()
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        if not s:
            flush()
            i += 1
            continue
        if s in ("[[TOC]]", "[[LOT]]", "[[LOF]]", "[[LOC]]"):
            flush()
            instr = {"[[TOC]]": 'TOC \\o "1-3" \\h \\z \\u',
                     "[[LOT]]": 'TOC \\h \\z \\c "Tableau"',
                     "[[LOF]]": 'TOC \\h \\z \\c "Figure"',
                     "[[LOC]]": 'TOC \\h \\z \\c "Extrait"'}[s]
            b.field_block(instr, "Clic droit puis « Mettre à jour le champ » (F9).")
        elif s == "[[FIN_LIMINAIRES]]":
            flush()
            b.end_section(b.front_sect)
        elif s == "[[SAUT]]":
            flush()
            b.page_break()
        elif s.startswith("#= "):
            flush()
            b.front_title(s[3:].strip())
        elif s.startswith("#"):
            flush()
            m = re.match(r"^(#{1,3})(!?)\s+(.*)$", s)
            level, bang, text = len(m.group(1)), m.group(2), m.group(3)
            b.heading(text, level, numbered=not bang)
        elif s.startswith("Tableau:"):
            flush()
            caption = s[len("Tableau:"):].strip()
            i += 1
            rows = [split_row(lines[i])]
            aligns = parse_aligns(lines[i + 1])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            b.table(rows, aligns, caption)
            continue
        elif s.startswith("Figure:"):
            flush()
            parts = [x.strip() for x in s[len("Figure:"):].split("|")]
            path = ROOT / parts[1]
            if not path.exists():
                print(f"AVERTISSEMENT : figure absente {path}", file=sys.stderr)
            else:
                b.figure(parts[0], path, float(parts[2]),
                         landscape=len(parts) > 3 and parts[3] == "paysage")
        elif s.startswith("Capture:"):
            flush()
            slot, caption, rel, instruction = [x.strip() for x in s[len("Capture:"):].split("|")]
            b.capture(caption, rel, instruction, slot)
        elif s.startswith("Code:"):
            flush()
            parts = [x.strip() for x in s[len("Code:"):].split("|")]
            slot, caption, spec = parts[:3]
            b.code(caption, spec, parts[3] if len(parts) > 3 else "", slot)
        elif s.startswith(">"):
            flush()
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            b.note(" ".join(buf))
            continue
        elif s.startswith("- "):
            flush()
            item = [s[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() \
                    and not lines[i].strip().startswith("- "):
                item.append(lines[i].strip())
                i += 1
            b.list_item(" ".join(item))
            continue
        elif re.match(r"^\d+\.\s", s):
            flush()
            if ordered_id is None:
                ordered_id = b._new_ordered_num()
            item = [re.sub(r"^\d+\.\s", "", s)]
            i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
                item.append(lines[i].strip())
                i += 1
            b.list_item(" ".join(item), ordered_num_id=ordered_id)
            continue
        else:
            para_buf.append(s)
        i += 1
    flush()
    b.finish(out)
    return b


def update_with_word(out: Path, pdf: Path | None):
    """Met à jour sommaire, listes et champs SEQ via Word (COM), puis enregistre."""
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
try {{
  $d = $w.Documents.Open('{out}')
  $d.Fields.Update() | Out-Null
  foreach ($t in $d.TablesOfContents) {{ $t.Update() | Out-Null }}
  $d.Repaginate()
  foreach ($t in $d.TablesOfContents) {{ $t.Update() | Out-Null }}
  $d.Save()
  {"$d.ExportAsFixedFormat('" + str(pdf) + "', 17)" if pdf else ""}
  "pages: " + $d.ComputeStatistics(2)
  "mots: " + $d.ComputeStatistics(0)
  $d.Close()
}} finally {{ $w.Quit() }}
"""
    res = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                         capture_output=True, text=True)
    print(res.stdout.strip())
    if res.returncode != 0:
        print(res.stderr.strip(), file=sys.stderr)
        print("Word indisponible : ouvrir le document et appuyer sur F9 pour les champs.",
              file=sys.stderr)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--source", type=Path, default=SOURCE)
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--pdf", type=Path, default=None, help="exporte aussi un PDF (via Word)")
    ap.add_argument("--no-word", action="store_true", help="ne pas mettre à jour les champs via Word")
    args = ap.parse_args(argv)
    b = build(args.source, args.out)
    print(f"{args.out} : {b.counters['Tableau']} tableaux, {b.counters['Figure']} figures, "
          f"{b.counters['Extrait']} extraits de code")
    if b.slots:
        print("Emplacements (documents/captures/) :")
        for kind, slot, rel, found in b.slots:
            if kind == "capture":
                state = "inséré" if found else "EN ATTENTE"
            else:
                state = "capture insérée" if found else "code du dépôt (capture facultative)"
            print(f"  {slot:<4} {state:<36} {rel}")
    if not args.no_word:
        update_with_word(args.out.resolve(), args.pdf.resolve() if args.pdf else None)


if __name__ == "__main__":
    main()
