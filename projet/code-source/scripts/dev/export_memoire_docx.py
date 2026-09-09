"""Exporte les chapitres du mémoire (Markdown) vers un document Word (.docx).

Usage :
    python scripts/dev/export_memoire_docx.py [--out documents/memoire_M2_MBDS.docx]

Convertit `chapters/01..06.md` en un unique .docx : titres, paragraphes, listes,
tableaux, blocs de code et citations. Les diagrammes Mermaid restent en bloc de code
(python-docx ne les rend pas en image).
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[4]
CHAPTERS_DIR = ROOT / "chapters"


def add_runs(paragraph, text: str) -> None:
    """Applique **gras** et `code` inline dans un paragraphe."""
    pattern = re.compile(r"(\*\*.*?\*\*|`[^`]*`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(1)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table(lines: list[str]) -> str:
    """Supprime la ligne de séparation éventuelle d'un tableau markdown."""
    if len(lines) >= 2 and re.match(r"^\|[\s:\-|]*\|$", lines[1]):
        return "\n".join([lines[0]] + lines[2:])
    return "\n".join(lines)


def add_markdown(doc: Document, md: str) -> None:
    lines = md.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            code_start = i
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                i += 1
            code_lines = lines[code_start + 1 : i]
            for j, cl in enumerate(code_lines):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(cl if j < len(code_lines) - 1 else cl)
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            if i >= n:
                break
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            table_md = parse_table(table_lines)
            t_rows = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in table_md.splitlines()
            ]
            if not t_rows:
                continue
            ncols = max(len(r) for r in t_rows)
            table = doc.add_table(rows=len(t_rows), cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate(t_rows):
                for c_idx in range(ncols):
                    cell = row[c_idx] if c_idx < len(row) else ""
                    para = table.cell(r_idx, c_idx).paragraphs[0]
                    add_runs(para, cell)
                    if r_idx == 0:
                        for run in para.runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_runs(p, stripped[2:])
            for run in p.runs:
                if run.text.strip():
                    run.italic = True
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph()
            add_runs(p, stripped)
            p.paragraph_format.left_indent = Pt(18)
        elif stripped.startswith("- "):
            p = doc.add_paragraph()
            add_runs(p, "• " + stripped[2:])
            p.paragraph_format.left_indent = Pt(18)
        elif stripped == "---":
            p = doc.add_paragraph()
            p.add_run("―" * 40)
            p.paragraph_format.space_after = Pt(6)
        elif stripped == "":
            pass
        else:
            p = doc.add_paragraph()
            add_runs(p, stripped)
        i += 1


def build(out_path: Path) -> None:
    doc = Document()

    title = doc.add_heading("Mémoire de stage — Master MBDS", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        "Plateforme Big Data de gestion et de gouvernance des données patients : "
        "nettoyage, déduplication et contrôle d'accès basé sur le consentement"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    for chapter in sorted(CHAPTERS_DIR.glob("0*.md")):
        add_markdown(doc, chapter.read_text(encoding="utf-8"))
        doc.add_page_break()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"DOCX écrit : {out_path.resolve()}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Exporte les chapitres vers un .docx.")
    parser.add_argument(
        "--out",
        default=str(ROOT / "documents" / "memoire_M2_MBDS.docx"),
        help="Chemin de sortie du fichier .docx",
    )
    args = parser.parse_args()
    build(Path(args.out))


if __name__ == "__main__":
    main()